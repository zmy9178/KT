from __future__ import annotations

import argparse
import json
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

from .defaults import DEFAULT_JSS_P, DEFAULT_ZHS_P
from .simulation import run_exit_stages, run_full_simulation


KT_DATA_DIR = Path.home() / "AppData" / "Roaming" / "KTData"

FIRST_INPUT_NAME = "airdrop_input_first.xlsx"
SECOND_INPUT_NAME = "airdrop_input_second.xlsx"
FIRST_OUTPUT_NAME = "airdrop_output_first.xlsx"
SECOND_OUTPUT_NAME = "airdrop_output_second.xlsx"
STATUS_NAME = "airdrop_status.json"

FIRST_OUTPUT_COLUMNS = (
    "x_position",
    "y_position",
    "x_velocity",
    "y_velocity",
    "pitch_angle",
    "roll_angle",
)

SECOND_OUTPUT_COLUMNS = (
    "x_position",
    "y_position",
    "z_position",
    "x_velocity",
    "y_velocity",
    "z_velocity",
    "pitch_angle",
    "roll_angle",
    "yaw_angle",
)

PARAMETER_LENGTHS = {
    "kte_p": 9,
    "vel_p": 19,
    "qys_p": 4,
    "jss_p": 11,
    "zhs_p": (19, 20),
    "sim_p": 2,
}


def read_parameter_file(
    path: Path,
    required_names: Iterable[str],
) -> dict[str, list[float]]:
    """Read row-based parameter vectors from Sheet1."""

    if not path.exists():
        raise FileNotFoundError(f"input file does not exist: {path}")

    required = tuple(required_names)
    expected = {name: PARAMETER_LENGTHS[name] for name in required}
    found: dict[str, list[float]] = {}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        for row in worksheet.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue

            name = str(row[0]).strip()
            if name not in expected:
                continue

            values: list[float] = []
            for cell in row[1:]:
                if cell is None:
                    continue
                if isinstance(cell, str) and not cell.strip():
                    continue
                values.append(float(cell))

            lengths = (
                (expected[name],)
                if isinstance(expected[name], int)
                else tuple(expected[name])
            )
            if len(values) not in lengths:
                expected_text = " or ".join(str(length) for length in lengths)
                raise ValueError(
                    f"{path.name} row {name} must contain {expected_text} values; "
                    f"got {len(values)}"
                )
            found[name] = values
    finally:
        workbook.close()

    missing = [name for name in required if name not in found]
    if missing:
        raise ValueError(f"{path.name} missing required rows: {', '.join(missing)}")
    return found


def run_first_stage_from_files(data_dir: Path = KT_DATA_DIR) -> Path:
    """Run the exit stage and write the six-column first-stage Excel output."""

    first_input = read_parameter_file(
        data_dir / FIRST_INPUT_NAME,
        ("vel_p", "qys_p", "sim_p"),
    )
    second_input = read_parameter_file(
        data_dir / SECOND_INPUT_NAME,
        ("kte_p",),
    )

    result = run_exit_stages(
        second_input["kte_p"],
        first_input["vel_p"],
        first_input["qys_p"],
        DEFAULT_JSS_P,
        DEFAULT_ZHS_P,
        first_input["sim_p"],
    )

    output = _first_stage_output(result)
    output_path = data_dir / FIRST_OUTPUT_NAME
    write_xlsx(output_path, FIRST_OUTPUT_COLUMNS, output)
    return output_path


def run_second_stage_from_files(data_dir: Path = KT_DATA_DIR) -> Path:
    """Run the full airdrop simulation and write the nine-column Excel output."""

    first_input = read_parameter_file(
        data_dir / FIRST_INPUT_NAME,
        ("vel_p", "qys_p", "sim_p"),
    )
    second_input = read_parameter_file(
        data_dir / SECOND_INPUT_NAME,
        ("kte_p", "jss_p", "zhs_p"),
    )

    result = run_full_simulation(
        second_input["kte_p"],
        first_input["vel_p"],
        first_input["qys_p"],
        second_input["jss_p"],
        second_input["zhs_p"],
        first_input["sim_p"],
    )
    output = _second_stage_output(np.asarray(result["data_all"], dtype=float))

    output_path = data_dir / SECOND_OUTPUT_NAME
    write_xlsx(output_path, SECOND_OUTPUT_COLUMNS, output)
    return output_path


def _first_stage_output(result: dict[str, object]) -> np.ndarray:
    ctx = result["ctx"]
    time_1 = np.asarray(result["time_1"], dtype=float)
    data_1 = np.asarray(result["data_1"], dtype=float)
    time_2 = np.asarray(result["time_2"], dtype=float)
    data_2 = np.asarray(result["data_2"], dtype=float)
    time_pd = float(result["time_pd"])
    time_fz = float(result["time_fz"])

    exit_stage_1 = _take_stage_prefix(time_1, data_1, time_pd)
    exit_stage_2 = _take_stage_prefix(time_2, data_2, time_fz)

    data_cc_1 = np.zeros((len(exit_stage_1), 6), dtype=float)
    data_cc_1[:, :4] = exit_stage_1
    data_cc = np.vstack((data_cc_1, exit_stage_2[1:]))

    return np.column_stack(
        (
            -data_cc[:, 0],
            data_cc[:, 2],
            ctx.drop_envi.airplane_velo / 3.6 + data_cc[:, 1],
            data_cc[:, 3],
            np.rad2deg(data_cc[:, 4]),
            np.zeros(len(data_cc), dtype=float),
        )
    )


def _second_stage_output(data_all: np.ndarray) -> np.ndarray:
    if data_all.ndim != 2 or data_all.shape[1] != 9:
        raise ValueError(
            "full simulation output must be a two-dimensional array with 9 columns"
        )

    output = data_all.copy()
    # The migrated dynamics do not calculate yaw. The old ninth column is
    # overload, so expose yaw as 0 instead of relabeling overload as an angle.
    output[:, 8] = 0.0
    return output


def _take_stage_prefix(
    time_values: np.ndarray,
    data_values: np.ndarray,
    stop_time: float,
) -> np.ndarray:
    stop_index = int(np.searchsorted(time_values, stop_time, side="left"))
    stop_index = min(max(stop_index, 0), len(data_values) - 1)
    return data_values[: stop_index + 1]


def write_xlsx(path: Path, columns: Iterable[str], data: np.ndarray) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(f"{path.stem}.tmp{path.suffix}")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(list(columns))
    for row in np.asarray(data, dtype=float):
        worksheet.append([float(value) for value in row])
    workbook.save(temp_path)
    temp_path.replace(path)


def write_status(
    data_dir: Path,
    *,
    stage: str,
    status: str,
    message: str,
) -> None:
    status_path = data_dir / STATUS_NAME
    payload = {
        "stage": stage,
        "status": status,
        "message": message,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    status_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = status_path.with_name(f"{status_path.stem}.tmp{status_path.suffix}")
    temp_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temp_path.replace(status_path)


def _run_with_status(stage: str, data_dir: Path) -> list[Path]:
    written: list[Path] = []
    write_status(data_dir, stage=stage, status="running", message="simulation running")
    try:
        if stage in ("first", "all"):
            written.append(run_first_stage_from_files(data_dir))
        if stage in ("second", "all"):
            written.append(run_second_stage_from_files(data_dir))
    except Exception as exc:
        write_status(data_dir, stage=stage, status="error", message=str(exc))
        raise

    message = "; ".join(str(path) for path in written)
    write_status(data_dir, stage=stage, status="success", message=message)
    return written


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Run airdrop Excel file bridge.")
    parser.add_argument(
        "stage",
        choices=("first", "second", "all"),
        nargs="?",
        default="all",
        help="which stage output to generate",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=KT_DATA_DIR,
        help="front-end exchange directory",
    )
    args = parser.parse_args(argv)

    for path in _run_with_status(args.stage, args.data_dir):
        print(f"Output saved to: {path}")


if __name__ == "__main__":
    main()
